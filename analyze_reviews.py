"""Разбор отзывов нейросетью GigaChat.

Для каждого отзыва определяет тон и темы, собирает таблицу Excel:
лист «Отзывы» — каждый отзыв с тоном и темами;
лист «Сводка» — сколько раз упоминается каждая тема и какая доля негатива (формулами);
лист «Главные жалобы» — три темы с самым большим числом негативных отзывов и примеры.

Берёт ключ и модель GigaChat из .env, поэтому клади файл в папку с ботом.

Запуск:
    pip install openpyxl
    python analyze_reviews.py отзывы_пример.csv
"""
import csv
import json
import os
import re
import sys
import time
from collections import Counter, defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

TOPICS = ["Размер и посадка", "Качество ткани и пошива", "Соответствие фото и описанию",
          "Цвет", "Доставка и упаковка", "Цена", "Другое"]
TONES = ["положительный", "нейтральный", "отрицательный"]
BATCH = 20

PROMPT = """Ты аналитик отзывов интернет-магазина одежды. Для каждого отзыва определи:
- tone: одно из значений "положительный", "нейтральный", "отрицательный";
- topics: список тем только из этого перечня: {topics}. Указывай только те темы, о которых в отзыве прямо говорится. Если ни одна не подходит, укажи ["Другое"].
Ответь только JSON-массивом без пояснений, по одному объекту на отзыв:
[{{"id": 1, "tone": "...", "topics": ["..."]}}]

Отзывы:
{items}"""

HEAD_FILL = PatternFill("solid", fgColor="BE185D")
HEAD_FONT = Font(name="Arial", bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial")


def make_llm():
    from dotenv import load_dotenv
    from gigachat import GigaChat
    from gigachat.models import Chat, Messages, MessagesRole

    load_dotenv()
    client = GigaChat(credentials=os.environ["GIGACHAT_CREDENTIALS"],
                      scope=os.getenv("GIGACHAT_SCOPE", "GIGACHAT_API_PERS"),
                      verify_ssl_certs=os.getenv("GIGACHAT_VERIFY_SSL", "false").lower() == "true")
    # Имена моделей у GigaChat меняются; если не задано, берём первую доступную по ключу.
    model = os.getenv("GIGACHAT_MODEL") or client.get_models().data[0].id_

    def ask(text):
        chat = Chat(model=model, messages=[Messages(role=MessagesRole.USER, content=text)], temperature=0.1)
        return client.chat(chat).choices[0].message.content
    return ask


def normalize(value, allowed):
    value = (value or "").strip().lower()
    for a in allowed:
        if value == a.lower() or value.startswith(a.lower()[:6]):
            return a
    return None


def parse_answer(answer):
    answer = re.sub(r"```(?:json)?", "", answer)
    start, end = answer.find("["), answer.rfind("]")
    data = json.loads(answer[start:end + 1])
    result = {}
    for item in data:
        tone = normalize(item.get("tone"), TONES)
        topics = [t for t in (normalize(x, TOPICS) for x in item.get("topics", [])) if t]
        result[int(item["id"])] = (tone, sorted(set(topics), key=TOPICS.index) or ["Другое"])
    return result


def classify(reviews, ask):
    labels = {}
    for start in range(0, len(reviews), BATCH):
        chunk = reviews[start:start + BATCH]
        items = "\n".join(f"{i + 1}. {r}" for i, r in enumerate(chunk))
        for attempt in range(3):
            try:
                parsed = parse_answer(ask(PROMPT.format(topics=", ".join(TOPICS), items=items)))
                break
            except Exception as e:
                print(f"  повтор из-за ошибки: {e}")
                time.sleep(2)
        else:
            parsed = {}
        for i in range(len(chunk)):
            labels[start + i] = parsed.get(i + 1, (None, ["Другое"]))
        print(f"Разобрано отзывов: {min(start + BATCH, len(reviews))} из {len(reviews)}")
    return labels


def style_header(ws, widths):
    for i, w in enumerate(widths, 1):
        c = ws.cell(row=1, column=i)
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def save(reviews, labels, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Отзывы"
    ws.append(["№", "Отзыв", "Тон", "Темы"])
    style_header(ws, [6, 90, 16, 48])
    for i, text in enumerate(reviews):
        tone, topics = labels[i]
        ws.append([i + 1, text, tone or "не распознан", ", ".join(topics)])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = BODY_FONT
            c.alignment = Alignment(wrap_text=c.column == 2, vertical="top")
    n = ws.max_row
    ws.auto_filter.ref = f"A1:D{n}"

    s = wb.create_sheet("Сводка")
    s.append(["Тема", "Упоминаний", "Из них отрицательных", "Доля негатива"])
    style_header(s, [32, 14, 22, 16])
    for r, topic in enumerate(TOPICS, start=2):
        s.append([topic,
                  f'=COUNTIF(\'Отзывы\'!D2:D{n},"*{topic}*")',
                  f'=COUNTIFS(\'Отзывы\'!D2:D{n},"*{topic}*",\'Отзывы\'!C2:C{n},"отрицательный")',
                  f"=IF(B{r}=0,0,C{r}/B{r})"])
        s.cell(row=r, column=4).number_format = "0%"
    base = len(TOPICS) + 3  # после тем одна пустая строка
    for k, tone in enumerate(TONES):
        s.cell(row=base + k, column=1, value=f"Отзывов: {tone}")
        s.cell(row=base + k, column=2, value=f'=COUNTIF(\'Отзывы\'!C2:C{n},"{tone}")')
    s.cell(row=base + 3, column=1, value="Всего отзывов")
    s.cell(row=base + 3, column=2, value=f"=SUM(B{base}:B{base + 2})")
    for row in s.iter_rows(min_row=2):
        for c in row:
            c.font = BODY_FONT

    neg_by_topic = Counter()
    examples, used = defaultdict(list), set()
    for i, text in enumerate(reviews):
        tone, topics = labels[i]
        if tone == "отрицательный":
            for t in topics:
                if t != "Другое":
                    neg_by_topic[t] += 1
                    if len(examples[t]) < 2 and i not in used:
                        examples[t].append(text if len(text) <= 160 else text[:157] + "...")
                        used.add(i)
    top = neg_by_topic.most_common(3)

    g = wb.create_sheet("Главные жалобы")
    g.append(["Тема", "Негативных отзывов", "Примеры"])
    style_header(g, [30, 20, 100])
    for topic, cnt in top:
        g.append([topic, cnt, "\n".join(f"«{e}»" for e in examples[topic])])
    for row in g.iter_rows(min_row=2):
        for c in row:
            c.font = BODY_FONT
            c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)
    return top


def main(csv_path, out="Разбор_отзывов.xlsx", ask=None):
    with open(csv_path, encoding="utf-8-sig") as f:
        reviews = [row["review"] for row in csv.DictReader(f) if row.get("review")]
    labels = classify(reviews, ask or make_llm())
    top = save(reviews, labels, out)
    tones = Counter(labels[i][0] for i in labels)
    print(f"\nГотово: {out}")
    print("Тон:", ", ".join(f"{t} — {tones.get(t, 0)}" for t in TONES))
    print("Главные жалобы:", "; ".join(f"{t} ({c})" for t, c in top))


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    main(sys.argv[1])
